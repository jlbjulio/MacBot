"""Validated design choices shared by generated documents and presentations."""
import re
from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import Literal, cast

from pydantic import BaseModel, Field, model_validator

FontName = Literal["Calibri", "Arial", "Georgia", "Aptos", "Verdana", "Segoe UI", "Cambria", "Trebuchet MS", "Tahoma", "Times New Roman"]
FONT_FILES = {
    "Calibri": ("calibri", "calibrib", "calibrii", "calibriz"),
    "Arial": ("arial", "arialbd", "ariali", "arialbi"),
    "Georgia": ("georgia", "georgiab", "georgiai", "georgiaz"),
    "Aptos": ("aptos", "aptos-bold", "aptos-italic", "aptos-bold-italic"),
    "Verdana": ("verdana", "verdanab", "verdanai", "verdanaz"),
    "Segoe UI": ("segoeui", "segoeuib", "segoeuii", "segoeuiz"),
    "Cambria": ("cambria", "cambriab", "cambriai", "cambriaz"),
    "Trebuchet MS": ("trebuc", "trebucbd", "trebucit", "trebucbi"),
    "Tahoma": ("tahoma", "tahomabd", "tahoma", "tahomabd"),
    "Times New Roman": ("times", "timesbd", "timesi", "timesbi"),
}


def luminance(color):
    channels = [int(color[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    return sum((c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055)**2.4) * weight
               for c, weight in zip(channels, (0.2126, 0.7152, 0.0722)))


def readable_color(foreground, background):
    first, second = sorted((luminance(background), luminance(foreground)))
    if (second + 0.05) / (first + 0.05) >= 4.5:
        return foreground.upper()
    return "151515" if luminance(background) > 0.179 else "FAFAFA"


def font_file(name, variant=0):
    import os
    roots = [Path("C:/Windows/Fonts"), Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft/Windows/Fonts"]
    stem = FONT_FILES[name][variant]
    return next((root / (stem + suffix) for root in roots for suffix in (".ttf", ".ttc")
                 if (root / (stem + suffix)).is_file()), None)


def pdf_font(name):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    selected = name if font_file(name) else "Arial"
    regular = font_file(selected)
    if regular is None:
        return "Times-Roman" if name in ("Georgia", "Cambria", "Times New Roman") else "Helvetica"
    family = "MacBot" + re.sub(r"\W", "", selected)
    variants = []
    for number in range(4):
        path = font_file(selected, number) or regular
        identifier = family + str(number)
        if identifier not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(identifier, str(path)))
        variants.append(identifier)
    pdfmetrics.registerFontFamily(variants[0], normal=variants[0], bold=variants[1], italic=variants[2], boldItalic=variants[3])
    return variants[0]


class Theme(BaseModel):
    background: str = Field(default="F5F1E8", pattern=r"^[0-9A-Fa-f]{6}$")
    foreground: str = Field(default="203430", pattern=r"^[0-9A-Fa-f]{6}$")
    accent: str = Field(default="B46B47", pattern=r"^[0-9A-Fa-f]{6}$")
    font: FontName = "Calibri"
    heading_font: FontName | None = None
    style: Literal["editorial", "minimal", "formal", "bold"] = "editorial"
    body_size: float = Field(default=11, ge=9, le=28)
    heading_size: float = Field(default=20, ge=14, le=48)
    title_size: float = Field(default=34, ge=24, le=72)
    line_spacing: float = Field(default=1.15, ge=1.05, le=1.6)
    spacing: Literal["compact", "comfortable", "airy"] = "comfortable"

    @property
    def heading_face(self):
        return self.heading_font or self.font

    @property
    def paragraph_gap(self):
        return {"compact": 5, "comfortable": 9, "airy": 14}[self.spacing]

    @model_validator(mode="after")
    def readable(self):
        self.foreground = readable_color(self.foreground, self.background)
        return self


def requested_theme(value, request):
    theme = Theme.model_validate(value)
    palette = {"purple": "674183", "blue": "285F9C", "green": "28634A", "teal": "236B70",
               "red": "A93636", "orange": "B8682D", "pink": "BA5B86", "yellow": "C89D28",
               "navy": "182F50", "cream": "F5F1E8", "beige": "EDE2CD", "white": "FFFFFF",
               "black": "151515", "gray": "777777", "grey": "777777"}
    if re.search(r"\b(?:theme|palette|colou?rs?|background|accent)\b", request, re.I):
        names = re.findall(r"\b(?:" + "|".join(palette) + r")\b", request, re.I)
        neutral = {"cream", "beige", "white", "black"}
        background = next((name.lower() for name in names if name.lower() in neutral), None)
        accent = next((name.lower() for name in names if name.lower() not in neutral), None)
        if background:
            theme.background = palette[background]
        if accent:
            theme.accent = palette[accent]
    for role in ("background", "foreground", "accent"):
        explicit = re.search(r"\b" + role + r"\s*(?:colou?r\s*)?(?:[:=]|is)?\s*#([0-9a-f]{6})\b", request, re.I)
        if explicit:
            setattr(theme, role, explicit[1].upper())
    faces = "|".join(re.escape(name) for name in FONT_FILES)
    specific = re.search(r"\b(?:headings?|titles?)\s+(?:in|font\s*[:=]?)\s*(" + faces + r")\b", request, re.I)
    fonts = re.findall(r"\b(?:" + faces + r")\b", request, re.I)
    if specific:
        theme.heading_font = next(name for name in FONT_FILES if name.lower() == specific[1].lower())
    body = re.search(r"\b(?:body|text)\s+(?:in|font\s*[:=]?)\s*(" + faces + r")\b", request, re.I)
    if body:
        theme.font = next(name for name in FONT_FILES if name.lower() == body[1].lower())
    elif len(set(face.lower() for face in fonts)) == 1 and not specific:
        theme.font = next(name for name in FONT_FILES if name.lower() == fonts[0].lower())
        theme.heading_font = theme.font
    for role in ("body", "heading", "title"):
        size = re.search(r"\b" + role + r"(?:\s+font)?\s+(?:size\s*[:=]?\s*)?(\d+(?:\.\d+)?)\s*(?:pt|points?)\b", request, re.I)
        if size:
            setattr(theme, role + "_size", float(size[1]))
    return Theme.model_validate(theme.model_dump()).model_dump()


class Table(BaseModel):
    columns: list[str] = Field(min_length=1, max_length=8)
    rows: list[list[str | int | float | None]] = Field(max_length=40)


class Section(BaseModel):
    heading: str = Field(max_length=160)
    text: str = Field(default="", max_length=3500)
    bullets: list[str] = Field(default_factory=list, max_length=12)
    table: Table | None = None
    emphasis: Literal["normal", "callout"] = "normal"
    page_break_before: bool = False


class DocumentSpec(BaseModel):
    title: str = Field(max_length=160)
    subtitle: str = Field(default="", max_length=250)
    theme: Theme = Field(default_factory=Theme)
    format: Literal["docx", "pdf", "both"] = "docx"
    paper: Literal["A4", "LETTER"] = "A4"
    orientation: Literal["portrait", "landscape"] = "portrait"
    cover: bool = True
    margin: float = Field(default=0.85, ge=0.45, le=1.3)
    sections: list[Section] = Field(min_length=1, max_length=20)


@dataclass
class DocumentBlock:
    kind: Literal["paragraph", "bullet", "number", "heading", "table"]
    text: str = ""
    table: Table | None = None


def document_blocks(text: str) -> list[DocumentBlock]:
    """Keep model-supplied Markdown usable when it lands in a text field."""
    blocks: list[DocumentBlock] = []
    lines = text.splitlines()
    paragraph: list[str] = []

    def flush():
        if paragraph:
            blocks.append(DocumentBlock("paragraph", "\n".join(paragraph)))
            paragraph.clear()

    def cells(line: str) -> list[str]:
        return [cell.strip() for cell in line.strip().strip("|").split("|")]

    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if ("|" in line and index + 1 < len(lines)
                and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells(lines[index + 1]))
                and len(cells(line)) == len(cells(lines[index + 1]))):
            flush()
            columns, rows = cells(line), []
            index += 2
            while index < len(lines) and "|" in lines[index]:
                row = cells(lines[index])
                if len(row) != len(columns):
                    break
                rows.append(row)
                index += 1
            blocks.append(DocumentBlock("table", table=Table(columns=columns, rows=rows)))
            continue
        bullet = re.match(r"^[-*+]\s+(.+)", line)
        number = re.match(r"^\d+[.)]\s+(.+)", line)
        heading = re.match(r"^#{1,6}\s+(.+)", line)
        if not line or bullet or number or heading:
            flush()
            if bullet:
                blocks.append(DocumentBlock("bullet", bullet[1]))
            elif number:
                blocks.append(DocumentBlock("number", number[1]))
            elif heading:
                blocks.append(DocumentBlock("heading", heading[1]))
        else:
            paragraph.append(line)
        index += 1
    flush()
    return blocks


def prepare_document(value, request):
    spec = DocumentSpec.model_validate(value)
    if re.search(r"\bletter\s+(?:paper|page|size)|\b(?:paper|page)\s+(?:size\s+)?letter\b", request, re.I):
        spec.paper = "LETTER"
    if re.search(r"\b(?:no|without)\s+(?:a\s+)?(?:cover|title page)\b", request, re.I):
        spec.cover = False
    intent = re.sub(r"\btable of contents\b", "", request, flags=re.I)
    if (not re.search(r"\b(?:tables?|tablas?)\b", intent, re.I)
            or re.search(r"\b(?:no|without|avoid|sin)\s+(?:any\s+)?(?:tables?|tablas?)\b", intent, re.I)):
        return spec.model_dump()
    if any(section.table or any(block.kind == "table" for block in document_blocks(section.text))
           for section in spec.sections):
        return spec.model_dump()
    timeline = bool(re.search(r"\b(?:timeline|schedule|cronograma|calendario)\b", intent, re.I))
    section = next((section for section in spec.sections
                    if timeline and re.search(r"timeline|schedule|cronograma|calendario", section.heading, re.I)),
                   next((section for section in reversed(spec.sections) if section.bullets), spec.sections[-1]))
    if section.bullets:
        rows = []
        for index, bullet in enumerate(section.bullets, 1):
            label, separator, detail = bullet.partition(":")
            rows.append([label.strip(), detail.strip()] if separator else [str(index), bullet])
        section.table = Table(columns=["Period" if timeline else "Item", "Details"], rows=rows)
        section.bullets = []
    else:
        # Reformat planned content without inventing dates, figures or extra facts.
        section.table = Table(columns=["Section", "Details"], rows=[[section.heading, section.text]])
        section.text = ""
    return spec.model_dump()


def inline_parts(text: str):
    for part in re.split(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        bold = part.startswith(("**", "__")) and len(part) > 4
        italic = part.startswith("*") and not bold and len(part) > 2
        code = part.startswith("`") and len(part) > 2
        yield part[2:-2] if bold else part[1:-1] if italic or code else part, bold, italic


def pdf_inline(text: str) -> str:
    return "".join(("<b>" + escape(part) + "</b>" if bold else
                    "<i>" + escape(part) + "</i>" if italic else escape(part))
                   for part, bold, italic in inline_parts(text)).replace("\n", "<br/>")


def shade(cell, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(element)


def write_document(value, path, fmt):
    from docx import Document
    from docx.styles.style import ParagraphStyle as WordParagraphStyle
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    spec = DocumentSpec.model_validate(value)
    cover_sections = [section for section in spec.sections if spec.cover and len(spec.sections) > 1
                      and section.heading.strip().lower() in {"title page", "cover", "cover page"}]
    sections = [section for section in spec.sections if section not in cover_sections]

    def section_blocks(section: Section):
        blocks = document_blocks(section.text)
        blocks.extend(DocumentBlock("bullet", bullet) for bullet in section.bullets)
        if section.table:
            blocks.append(DocumentBlock("table", table=section.table))
        return blocks

    if fmt == "docx":
        doc = Document()
        page = doc.sections[0]
        width, height = (8.27, 11.69) if spec.paper == "A4" else (8.5, 11)
        if spec.orientation == "landscape":
            width, height = height, width
        page.page_width, page.page_height = Inches(width), Inches(height)
        page.top_margin = page.bottom_margin = Inches(spec.margin)
        page.left_margin = page.right_margin = Inches(spec.margin)
        def set_font(style, face):
            style.font.name = face
            fonts = style.element.xpath("./w:rPr/w:rFonts")[0]
            fonts.attrib.pop(qn("w:cstheme"), None)
            for role in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.attrib.pop(qn("w:" + role + "Theme"), None)
                fonts.set(qn("w:" + role), face)
        normal = cast(WordParagraphStyle, doc.styles["Normal"])
        set_font(normal, spec.theme.font)
        normal.font.size = Pt(spec.theme.body_size)
        normal.font.color.rgb = RGBColor.from_string(readable_color(spec.theme.foreground, "FFFFFF"))
        heading_color = readable_color(spec.theme.accent, "FFFFFF")
        normal.paragraph_format.space_after = Pt(spec.theme.paragraph_gap)
        normal.paragraph_format.line_spacing = spec.theme.line_spacing
        for heading in ("Title", "Heading 1", "Heading 2"):
            style = cast(WordParagraphStyle, doc.styles[heading])
            set_font(style, spec.theme.heading_face)
            style.font.size = Pt(spec.theme.title_size if heading == "Title" else
                                 spec.theme.heading_size if heading == "Heading 1" else max(14, spec.theme.heading_size - 4))
            style.font.bold = spec.theme.style != "minimal"
            style.font.color.rgb = RGBColor.from_string(heading_color)
            style.paragraph_format.keep_with_next = True
        for name in ("Header", "Footer", "Subtitle"):
            style = cast(WordParagraphStyle, doc.styles[name])
            set_font(style, spec.theme.font)
            style.font.color.rgb = normal.font.color.rgb
        doc.core_properties.title, doc.core_properties.author = spec.title, "MacBot"
        title_paragraph = doc.add_heading(spec.title, 0)
        if spec.theme.style in ("bold", "editorial"):
            border = OxmlElement("w:pBdr")
            edge = OxmlElement("w:bottom")
            for key, setting in {"val": "single", "sz": "24" if spec.theme.style == "bold" else "8",
                                 "space": "12", "color": spec.theme.accent}.items():
                edge.set(qn("w:" + key), setting)
            border.append(edge)
            title_paragraph._p.get_or_add_pPr().append(border)
        if spec.theme.style == "bold":
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), spec.theme.background)
            title_paragraph._p.get_or_add_pPr().append(fill)
            for run in title_paragraph.runs:
                run.font.color.rgb = RGBColor.from_string(spec.theme.foreground)
        if spec.subtitle:
            doc.add_paragraph(spec.subtitle, "Subtitle")
        def word_text(paragraph, text):
            for part, bold, italic in inline_parts(text):
                run = paragraph.add_run(part)
                run.bold, run.italic = bold, italic

        if spec.cover:
            for section in cover_sections:
                word_text(doc.add_paragraph(), section.text)
            if len(sections) > 2:
                doc.add_heading("Inside", 2)
                for section in sections:
                    doc.add_paragraph(section.heading, "List Bullet")
            doc.add_page_break()
        page.header.paragraphs[0].text = spec.title[:90]
        footer_paragraph = page.footer.paragraphs[0]
        footer_paragraph.text = "Page "
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer_paragraph._p.append(field)
        for section in sections:
            if section.page_break_before and section is not sections[0]:
                doc.add_page_break()
            section_heading = doc.add_heading(section.heading, 1)
            if section.emphasis == "callout":
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), spec.theme.background)
                section_heading._p.get_or_add_pPr().append(fill)
                for run in section_heading.runs:
                    run.font.color.rgb = RGBColor.from_string(spec.theme.foreground)
            for block in section_blocks(section):
                if block.kind != "table":
                    style = {"bullet": "List Bullet", "number": "List Number", "heading": "Heading 2"}.get(block.kind)
                    word_text(doc.add_paragraph(style=style), block.text)
                    continue
                assert block.table is not None
                data = block.table
                table = doc.add_table(rows=1, cols=len(data.columns))
                table.style = "Table Grid"
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for cell, heading in zip(table.rows[0].cells, data.columns):
                    word_text(cell.paragraphs[0], heading)
                    shade(cell, spec.theme.accent)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string(Theme(background=spec.theme.accent, foreground="FFFFFF").foreground)
                for number, row in enumerate(data.rows):
                    if len(row) != len(data.columns):
                        raise ValueError("Document table rows must match their columns.")
                    for cell, item in zip(table.add_row().cells, row):
                        word_text(cell.paragraphs[0], str(item) if item is not None else "")
                        if number % 2 == 0:
                            shade(cell, spec.theme.background)
                            for run in cell.paragraphs[0].runs:
                                run.font.color.rgb = RGBColor.from_string(spec.theme.foreground)
                doc.add_paragraph()
        doc.save(path)
        Document(path)
        return
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, LETTER, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table as PDFTable, TableStyle, Flowable
    font, heading_font = pdf_font(spec.theme.font), pdf_font(spec.theme.heading_face)
    page_size = A4 if spec.paper == "A4" else LETTER
    if spec.orientation == "landscape":
        page_size = landscape(page_size)
    accent = colors.HexColor("#" + spec.theme.accent)
    heading_color = colors.HexColor("#" + readable_color(spec.theme.accent, "FFFFFF"))
    body = ParagraphStyle("Body", fontName=font, fontSize=spec.theme.body_size,
                          leading=spec.theme.body_size * spec.theme.line_spacing,
                          spaceAfter=spec.theme.paragraph_gap,
                          textColor=colors.HexColor("#" + readable_color(spec.theme.foreground, "FFFFFF")))
    title = ParagraphStyle("Title", parent=body, fontName=heading_font, fontSize=spec.theme.title_size,
                           leading=spec.theme.title_size * 1.2, textColor=heading_color, spaceAfter=20)
    heading = ParagraphStyle("Heading", parent=body, fontName=heading_font, fontSize=spec.theme.heading_size,
                             leading=spec.theme.heading_size * 1.25, textColor=heading_color,
                             spaceBefore=spec.theme.paragraph_gap + 8, spaceAfter=spec.theme.paragraph_gap,
                             keepWithNext=True)
    if spec.theme.style == "bold":
        title.backColor = colors.HexColor("#" + spec.theme.background)
        title.textColor = colors.HexColor("#" + spec.theme.foreground)
        title.borderPadding = 12
    elements: list[Flowable] = [Paragraph(escape(spec.title), title)]
    if spec.subtitle:
        elements.append(Paragraph(escape(spec.subtitle), body))
    if spec.cover:
        elements += [Paragraph(pdf_inline(section.text), body) for section in cover_sections]
        if len(sections) > 2:
            elements += [Paragraph("Inside", heading), *[Paragraph(escape(s.heading), body) for s in sections]]
        elements.append(PageBreak())
    for section in sections:
        if section.page_break_before and section is not sections[0]:
            elements.append(PageBreak())
        section_style = heading
        if section.emphasis == "callout":
            section_style = ParagraphStyle("Callout", parent=heading, borderPadding=8,
                backColor=colors.HexColor("#" + spec.theme.background),
                textColor=colors.HexColor("#" + spec.theme.foreground))
        elements.append(Paragraph(escape(section.heading), section_style))
        for block in section_blocks(section):
            if block.kind != "table":
                prefix = "• " if block.kind == "bullet" else "– " if block.kind == "number" else ""
                elements.append(Paragraph(prefix + pdf_inline(block.text), heading if block.kind == "heading" else body))
                continue
            assert block.table is not None
            data = block.table
            tinted = ParagraphStyle("TintedCell", parent=body, textColor=colors.HexColor("#" + spec.theme.foreground))
            rows = [[Paragraph(pdf_inline(str(c)), tinted) for c in data.columns]]
            for number, row in enumerate(data.rows):
                if len(row) != len(data.columns):
                    raise ValueError("Document table rows must match their columns.")
                rows.append([Paragraph(pdf_inline(str(c) if c is not None else ""), tinted if number % 2 else body) for c in row])
            margin = spec.margin * 72
            table = PDFTable(rows, colWidths=[(page_size[0] - margin * 2 - 12) / len(data.columns)] * len(data.columns), repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + spec.theme.background)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + spec.theme.background)]),
                ("LINEBELOW", (0, 0), (-1, 0), 1.2, accent), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, -1), 7)]))
            elements += [table, Spacer(1, 12)]
    def footer(canvas, document):
        margin = spec.margin * 72
        canvas.setStrokeColor(accent)
        canvas.line(margin, 42, page_size[0] - margin, 42)
        canvas.setFont(font, 8)
        canvas.setFillColor(body.textColor)
        canvas.drawString(margin, 28, spec.title[:65])
        canvas.drawRightString(page_size[0] - margin, 28, f"Page {document.page}")
    margin = spec.margin * 72
    SimpleDocTemplate(str(path), pagesize=page_size, rightMargin=margin, leftMargin=margin,
                      topMargin=margin, bottomMargin=max(margin, 60),
                      title=spec.title, author="MacBot").build(elements, onFirstPage=footer, onLaterPages=footer)
    from pypdf import PdfReader
    if not PdfReader(path).pages:
        raise ValueError("The generated document is empty.")


def document_from_markdown(text):
    sections, title, current = [], "Your MacBot conversation", {"heading": "Response", "text": "", "bullets": []}
    for line in text.splitlines():
        if line.startswith("# ") and not sections and not current["text"]:
            title = line[2:]
        elif re.match(r"^#{1,3}\s", line):
            if current["text"].strip() or current["bullets"]:
                sections.append(current)
            current = {"heading": re.sub(r"^#+\s", "", line), "text": "", "bullets": []}
        elif re.match(r"^\s*[-*]\s", line):
            current["bullets"].append(re.sub(r"^\s*[-*]\s", "", line))
        else:
            current["text"] += line + "\n"
    if current["text"].strip() or current["bullets"]:
        sections.append(current)
    expanded = []
    for section in sections:
        content = section["text"]
        for start in range(0, max(1, len(content)), 3500):
            expanded.append({"heading": section["heading"][:160] if not start else section["heading"][:145] + " (continued)",
                             "text": content[start:start + 3500], "bullets": section["bullets"][:12] if not start else []})
        for start in range(12, len(section["bullets"]), 12):
            expanded.append({"heading": section["heading"][:145] + " (continued)", "text": "", "bullets": section["bullets"][start:start + 12]})
    if len(expanded) > 20:
        raise ValueError("This response is too long for a styled export. Save it as Markdown or text.")
    return {"title": title[:160], "cover": False, "sections": expanded or [{"heading": "Response", "text": text[:3500]}]}
