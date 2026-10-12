import io
import zipfile
from pathlib import Path

MAX_EXTRACTED_CHARACTERS = 20_000_000


def append_text(output, text, separator="\n"):
    prefix = separator if output.tell() else ""
    if output.tell() + len(prefix) + len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("The extracted document exceeds the safe 20 million character limit. Split the file and retry; its content has not been truncated.")
    output.write(prefix)
    output.write(text)


def extract(name, content, directory=None):
    suffix = Path(name).suffix.lower()
    if suffix in (".docx", ".xlsx", ".pptx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50_000_000:
                raise ValueError("The expanded document exceeds the 50 MB limit.")
    if suffix in (".png", ".jpg", ".jpeg", ".webp"):
        signatures = (b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"RIFF")
        if not content.startswith(signatures):
            raise ValueError("This image has an invalid format.")
        from .media import normalized_image, ocr_image
        return "image", ocr_image(normalized_image(content), directory)
    if suffix in (".txt", ".md", ".csv", ".json", ".jsonl", ".py", ".ts", ".js"):
        output = io.StringIO()
        append_text(output, content.decode("utf-8-sig", errors="replace"))
        return "document", output.getvalue()
    if suffix == ".pdf":
        from contextlib import ExitStack
        from pypdf import PdfReader
        from .media import read_pdf_page
        from .ocr import DocumentOCR
        from .store import data_directory

        reader = PdfReader(io.BytesIO(content))
        output = io.StringIO()
        readable = False
        with ExitStack() as stack:
            ocr = None
            scanned = None
            for number, page in enumerate(reader.pages):
                text = page.extract_text(extraction_mode="layout", layout_mode_space_vertically=False,
                                         layout_mode_strip_rotated=False) or ""
                append_text(output, f"Page {number + 1}\n{text}", separator="\n\n")
                readable |= bool(text.strip())
                if not text.strip():
                    if ocr is None:
                        ocr = stack.enter_context(DocumentOCR(directory or data_directory()))
                    if scanned is None:
                        import pypdfium2 as pdfium
                        scanned = stack.enter_context(pdfium.PdfDocument(content))
                    text = read_pdf_page(scanned, number, ocr)
                    append_text(output, text)
                    readable |= bool(text.strip())
                    continue
                for picture in page.images:
                    if ocr is None:
                        ocr = stack.enter_context(DocumentOCR(directory or data_directory()))
                    extracted = ocr.read(picture.image)
                    if extracted and " ".join(extracted.split()) not in " ".join(text.split()):
                        append_text(output, extracted)
        if not readable:
            raise ValueError("No readable text was found in this document.")
        return "document", output.getvalue()
    if suffix == ".pptx":
        from pptx import Presentation
        from pptx.shapes.autoshape import Shape
        from pptx.shapes.graphfrm import GraphicFrame
        from pptx.shapes.picture import Picture
        from pptx.shapes.group import GroupShape
        slides = Presentation(io.BytesIO(content)).slides
        output = io.StringIO()
        from contextlib import ExitStack
        from .ocr import DocumentOCR
        from .media import normalized_image
        from .store import data_directory
        with ExitStack() as stack:
            ocr = None
            def read_shapes(shapes):
                nonlocal ocr
                for shape in shapes:
                    if isinstance(shape, GroupShape):
                        read_shapes(shape.shapes)
                    elif isinstance(shape, Shape) and shape.has_text_frame:
                        append_text(output, shape.text)
                    elif isinstance(shape, GraphicFrame) and shape.has_table:
                        for row in shape.table.rows:
                            append_text(output, " | ".join(cell.text for cell in row.cells))
                    elif isinstance(shape, GraphicFrame) and shape.has_chart:
                        chart = shape.chart
                        if chart.has_title and chart.chart_title.has_text_frame:
                            append_text(output, chart.chart_title.text_frame.text)
                        for category in chart.plots[0].categories:
                            append_text(output, str(category.label))
                        for series in chart.series:
                            append_text(output, f"{series.name}: " + " | ".join(str(v) for v in series.values))
                    elif isinstance(shape, Picture):
                        if ocr is None:
                            ocr = stack.enter_context(DocumentOCR(directory or data_directory()))
                        append_text(output, ocr.read(normalized_image(shape.image.blob)))
            for i, slide in enumerate(slides, 1):
                append_text(output, f"Slide {i}")
                read_shapes(slide.shapes)
                if slide.has_notes_slide:
                    append_text(output, slide.notes_slide.notes_text_frame.text)
        return "document", output.getvalue()
    if suffix == ".docx":
        from docx import Document
        from docx.table import Table

        doc = Document(io.BytesIO(content))
        output = io.StringIO()
        for item in doc.iter_inner_content():
            if isinstance(item, Table):
                for row in item.rows:
                    append_text(output, " | ".join(cell.text for cell in row.cells))
            else:
                append_text(output, item.text)
        images = [r.target_part for r in doc.part.rels.values() if not r.is_external and r.reltype.endswith("/image")]
        if images:
            from .ocr import DocumentOCR
            from .media import normalized_image
            from .store import data_directory
            with DocumentOCR(directory or data_directory()) as ocr:
                for image in images:
                    append_text(output, ocr.read(normalized_image(image.blob)))
        return "document", output.getvalue()
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        output = io.StringIO()
        try:
            for sheet in workbook:
                append_text(output, f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    append_text(output, " | ".join(str(cell) if cell is not None else "" for cell in row))
        finally:
            workbook.close()
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            images = [name for name in archive.namelist() if name.startswith("xl/media/")
                      and Path(name).suffix.lower() in (".png", ".jpg", ".jpeg", ".bmp", ".tiff")]
            if images:
                from .ocr import DocumentOCR
                from .media import normalized_image
                from .store import data_directory
                with DocumentOCR(directory or data_directory()) as ocr:
                    for image in images:
                        append_text(output, "Workbook image")
                        append_text(output, ocr.read(normalized_image(archive.read(image))))
        return "document", output.getvalue()
    raise ValueError("Unsupported file type. Use a PDF, document, spreadsheet, text, image, audio, or video file.")


def export_document(text, path, fmt):
    if fmt in ("docx", "pdf"):
        from .design import document_from_markdown, write_document
        write_document(document_from_markdown(text), path, fmt)
        return
    if fmt in ("md", "txt"):
        path.write_text(text, encoding="utf-8")
    else:
        raise ValueError("Unsupported export format.")
    if not path.exists() or path.stat().st_size == 0:
        raise ValueError("The generated file could not be verified.")
