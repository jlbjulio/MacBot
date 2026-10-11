from pathlib import Path


def render_deck(value, path, images=None):
    import io
    import math
    from PIL import Image, ImageFont, ImageOps
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt
    from .artifacts import Deck
    from .design import font_file, readable_color

    deck = Deck.model_validate(value)
    presentation = Presentation()
    width = 13.33 if deck.aspect_ratio == "16:9" else 10
    scale = width / 13.33
    presentation.slide_width, presentation.slide_height = Inches(width), Inches(7.5)
    presentation.core_properties.title, presentation.core_properties.author = deck.title, "MacBot"
    theme, images = deck.theme, images or {}
    foreground, background, accent = theme.foreground, theme.background, theme.accent

    def shape(slide, x, y, w, h, color, kind=MSO_SHAPE.RECTANGLE):
        item = slide.shapes.add_shape(kind, Inches(x * scale), Inches(y), Inches(w * scale), Inches(h))
        item.fill.solid()
        item.fill.fore_color.rgb = RGBColor.from_string(color)
        item.line.fill.background()
        return item

    def fill_text(frame, content, w, h, size, bold=False, color=None, heading=False):
        face = theme.heading_face if heading else theme.font
        font_path = font_file(face, 1 if bold else 0) or font_file(face)
        gap = min(10, theme.paragraph_gap)
        limit = max(8, int(size))
        paragraphs = content.split("\n")
        for candidate in range(limit, 7, -1):
            font = ImageFont.truetype(str(font_path), round(candidate * 4 / 3)) if font_path else ImageFont.load_default(size=round(candidate * 4 / 3))
            available = max(1, (w * scale * 72 - 4) * 4 / 3)
            lines = 0
            for line in paragraphs:
                line_width, count = 0.0, 1
                for word in line.split():
                    length = font.getlength(word + " ")
                    if length > available:
                        count += math.ceil(length / available) - 1 + bool(line_width)
                        line_width = length % available
                    elif line_width and line_width + length > available:
                        count += 1
                        line_width = length
                    else:
                        line_width += length
                lines += count
            if lines * candidate * 1.2 * theme.line_spacing + max(0, len(paragraphs) - 1) * gap <= h * 72:
                break
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left = frame.margin_right = Inches(0.02)
        frame.margin_top = frame.margin_bottom = 0
        for number, line in enumerate(paragraphs):
            paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name, paragraph.font.size = face, Pt(candidate)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color or foreground)
            paragraph.line_spacing = theme.line_spacing
            paragraph.space_after = Pt(gap if number < len(paragraphs) - 1 else 0)
        return frame

    def text(slide, content, x, y, w, h, size=22, bold=False, color=None, heading=False):
        frame = slide.shapes.add_textbox(Inches(x * scale), Inches(y), Inches(w * scale), Inches(h)).text_frame
        return fill_text(frame, content, w, h, size, bold, color, heading)

    def cards(slide, items, top):
        columns = 1 if len(items) <= 2 else 2
        rows = math.ceil(len(items) / columns)
        height = min(1.8, (6.5 - top) / max(1, rows))
        for index, item in enumerate(items):
            left, y = 0.85 + (index % columns) * 6.0, top + (index // columns) * height
            width = 11.6 if columns == 1 else 5.65
            if theme.style == "bold":
                shape(slide, left, y, width, height - 0.1, theme.accent)
                color = readable_color("FFFFFF", theme.accent)
            else:
                shape(slide, left, y + 0.04, 0.04, height - 0.16, accent)
                color = foreground
            text(slide, f"{index + 1:02d}", left + 0.12, y + 0.07, 0.55, 0.4, 18, True, color, True)
            text(slide, item, left + 0.83, y + 0.06, width - 1.0, height - 0.22,
                 max(18, theme.body_size), color=color)

    for number, data in enumerate(deck.slides, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = theme.accent if data.background == "accent" else theme.foreground if data.background == "contrast" else theme.background
        foreground = readable_color(theme.foreground if data.background == "base" else theme.background, background)
        accent = readable_color(theme.accent, background)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
        layout = "image" if data.image_id else data.layout
        if theme.style != "minimal":
            shape(slide, 0, 0, 0.12 if theme.style == "formal" else 0.2, 7.5, accent)
        text(slide, deck.title, 0.75, 0.38, 10.6, 0.35, size=10)
        text(slide, f"{number:02d}", 11.85, 0.35, 0.7, 0.4, size=14, bold=True, color=accent)
        shape(slide, 0.75, 6.82, 11.8, 0.025, accent)
        text(slide, deck.title, 0.75, 6.96, 11, 0.25, size=9)
        if layout in ("cover", "section"):
            if theme.style == "bold":
                shape(slide, 9.7, 1.5, 2.75, 4.8, theme.accent)
            elif theme.style == "editorial":
                text(slide, f"{number:02d}", 9.6, 2.0, 2.8, 2.0, size=100, color=accent, heading=True)
            elif theme.style == "formal":
                shape(slide, 0.85, 1.45, 4.3, 0.07, accent)
            text(slide, data.title, 0.82, 1.8, 8.4, 2.3,
                 size=max(34, theme.title_size), bold=theme.style != "minimal", heading=True)
            text(slide, "\n".join(filter(None, [data.subtitle or deck.subtitle, *data.bullets])),
                 0.85, 4.5, 8.1, 1.65, size=max(18, theme.body_size))
        else:
            text(slide, data.title, 0.8, 1.0, 11.6, 1.05,
                 size=max(28, theme.heading_size), bold=True, heading=True)
            if data.subtitle:
                text(slide, data.subtitle, 0.83, 2.08, 11.5, 0.55, size=16)
            top = 2.85 if data.subtitle else 2.55
            if layout == "image":
                text(slide, "\n".join(data.bullets), 0.9, top, 7.5, 6.4 - top, size=max(18, theme.body_size))
            elif layout == "table" or data.table:
                assert data.table is not None
                table = slide.shapes.add_table(len(data.table.rows) + 1, len(data.table.columns),
                    Inches(0.85 * scale), Inches(top), Inches(11.6 * scale), Inches(6.45 - top)).table
                rows = [data.table.columns, *data.table.rows]
                for row_number, row in enumerate(rows):
                    fill = theme.accent if row_number == 0 else background
                    for column_number, item in enumerate(row):
                        cell = table.cell(row_number, column_number)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(fill)
                        color = readable_color("FFFFFF", fill) if row_number == 0 else foreground
                        fill_text(cell.text_frame, str(item) if item is not None else "",
                                  11.6 / len(row) - 0.15, (6.45 - top) / len(rows) - 0.08,
                                  max(16, theme.body_size), row_number == 0, color, row_number == 0)
            elif layout == "metric":
                text(slide, data.metric_value or data.bullets[0], 0.85, top, 5.3, 1.6, size=78, bold=True, color=accent, heading=True)
                label = data.metric_label or (data.bullets[0] if data.bullets else "")
                text(slide, label, 0.9, top + 1.75, 5.4, 1.2, size=max(20, theme.body_size))
                details = data.bullets if data.metric_label else data.bullets[1:]
                text(slide, "\n".join(details), 7.0, top + 0.1, 5.1, 3.4, size=max(18, theme.body_size))
            elif layout == "quote":
                text(slide, "“", 0.78, top - 0.25, 1.0, 1.0, size=76, color=accent, heading=True)
                text(slide, "\n".join(data.bullets), 1.75, top + 0.05, 9.8, 3.45,
                     size=max(24, theme.heading_size), heading=True)
            elif layout == "timeline":
                entries = [entry.model_dump() for entry in data.columns] or [
                    {"heading": f"{i + 1:02d}", "text": item} for i, item in enumerate(data.bullets)]
                columns = min(4, len(entries))
                rows = math.ceil(len(entries) / columns)
                row_height = (6.45 - top) / rows
                for index, entry in enumerate(entries):
                    left = 0.9 + (index % columns) * (11.7 / columns)
                    y = top + (index // columns) * row_height
                    shape(slide, left, y + 0.02, 0.26, 0.26, accent, MSO_SHAPE.OVAL)
                    text(slide, entry["heading"], left, y + 0.42, 10.8 / columns, 0.6, 20, True, heading=True)
                    text(slide, entry["text"], left, y + 1.05, 10.8 / columns, row_height - 1.1,
                         size=max(16, theme.body_size))
            elif layout in ("split", "comparison") or data.columns:
                entries = [entry.model_dump() for entry in data.columns] or [
                    {"heading": "", "text": "\n".join(data.bullets[:len(data.bullets) // 2 or 1])},
                    {"heading": "", "text": "\n".join(data.bullets[len(data.bullets) // 2 or 1:])}]
                for index, entry in enumerate(entries):
                    left, width = 0.85 + index * (11.7 / len(entries)), 10.9 / len(entries)
                    shape(slide, left, top, width, 0.075, accent)
                    text(slide, entry["heading"], left + 0.08, top + 0.3, width - 0.15, 0.85, 23, True, heading=True)
                    text(slide, entry["text"], left + 0.08, top + 1.3, width - 0.2, 6.4 - top - 1.3,
                         size=max(16, theme.body_size))
            else:
                cards(slide, data.bullets, top)
        if data.image_id:
            image = images.get(data.image_id)
            if image is None or not Path(image).is_file():
                raise ValueError("A slide references an image that was not attached to this request.")
            with Image.open(image) as original:
                cropped = ImageOps.fit(original.convert("RGB"), (720, 720))
                buffer = io.BytesIO()
                cropped.save(buffer, format="PNG")
                buffer.seek(0)
                slide.shapes.add_picture(buffer, Inches(9.0 * scale), Inches(top),
                    width=Inches(3.25 * scale), height=Inches(min(3.5, 6.4 - top)))
        content = [*data.bullets, *[column.heading + ": " + column.text for column in data.columns]]
        if data.table:
            content += [" | ".join(str(cell) for cell in row) for row in [data.table.columns, *data.table.rows]]
        notes = data.notes + ("\n\nFull slide content:\n" + "\n".join(content) if content else "")
        if notes.strip():
            frame = slide.notes_slide.notes_text_frame
            assert frame is not None
            frame.text = notes.strip()
    presentation.save(path)
    assert len(Presentation(path).slides) == len(deck.slides)
