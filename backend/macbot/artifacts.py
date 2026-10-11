import gc
import json
import re
import wave
from typing import Annotated, Any, Literal, cast
from uuid import NAMESPACE_URL, uuid5

from pydantic import BaseModel, Field, model_validator

from .assets import VOICE_FILE, VOICE_PREFIX, prepare_asset
from .design import FONT_FILES, Table, Theme, DocumentSpec, readable_color, requested_theme, write_document


class Chart(BaseModel):
    kind: Literal["bar", "line", "pie"] = "bar"
    title: str = Field(default="", max_length=120)
    category_column: int = Field(default=1, ge=1, le=20)
    value_columns: list[int] = Field(default_factory=lambda: [2], min_length=1, max_length=5)


class ImageSpec(BaseModel):
    prompt: str = Field(min_length=1, max_length=1000)
    negative_prompt: str = Field(default="blurry, distorted, unreadable text", max_length=500)
    width: int = Field(default=512, ge=128, le=2048)
    height: int = Field(default=512, ge=128, le=2048)
    format: Literal["png", "jpeg", "webp"] = "png"
    style: Literal["none", "photographic", "anime", "watercolor", "cinematic", "pixel-art", "vintage", "illustration"] = "none"
    filter: Literal["none", "grayscale", "sepia", "vivid"] = "none"


CellFormat = Literal["text", "integer", "decimal", "percent", "usd", "eur", "date"]


class Sheet(BaseModel):
    title: str = Field(max_length=150)
    columns: list[str] = Field(min_length=1, max_length=20)
    rows: list[list[str | int | float | bool | None]] = Field(max_length=300)
    theme: Theme = Field(default_factory=Theme)
    formats: list[CellFormat] = Field(default_factory=list, max_length=20)
    charts: list[Chart] = Field(default_factory=list, max_length=3)
    notes: str = Field(default="", max_length=800)
    header_style: Literal["filled", "accent", "minimal"] = "filled"
    stripe_rows: bool = True
    column_widths: list[Annotated[float, Field(ge=8, le=70)]] = Field(default_factory=list, max_length=20)
    alignments: list[Literal["auto", "left", "center", "right"]] = Field(default_factory=list, max_length=20)
    chart_placement: Literal["beside", "below", "dashboard"] = "beside"
    chart_colors: list[Annotated[str, Field(pattern=r"^[0-9A-Fa-f]{6}$")]] = Field(default_factory=list, max_length=8)
    print_orientation: Literal["portrait", "landscape"] = "landscape"

    @model_validator(mode="after")
    def column_design(self):
        if self.column_widths and len(self.column_widths) != len(self.columns):
            raise ValueError("Column widths must match the spreadsheet columns.")
        if self.alignments and len(self.alignments) != len(self.columns):
            raise ValueError("Alignments must match the spreadsheet columns.")
        return self


class SlideColumn(BaseModel):
    heading: str = Field(min_length=1, max_length=80)
    text: str = Field(min_length=1, max_length=400)


class SlideTable(Table):
    columns: list[str] = Field(min_length=1, max_length=5)
    rows: list[list[str | int | float | None]] = Field(min_length=1, max_length=6)


class Slide(BaseModel):
    title: str = Field(max_length=150)
    bullets: list[str] = Field(max_length=8)
    layout: Literal["cover", "split", "comparison", "timeline", "quote", "metric", "bullets", "cards", "section", "table"] = "bullets"
    subtitle: str = Field(default="", max_length=200)
    columns: list[SlideColumn] = Field(default_factory=list, max_length=3)
    metric_value: str = Field(default="", max_length=30)
    metric_label: str = Field(default="", max_length=100)
    notes: str = Field(default="", max_length=1500)
    image_id: str = Field(default="", max_length=80)
    background: Literal["base", "accent", "contrast"] = "base"
    table: SlideTable | None = None

    @model_validator(mode="after")
    def has_content(self):
        if self.layout not in ("cover", "metric", "section") and not self.bullets and not self.columns and not self.table:
            raise ValueError("Content slides need meaningful bullets or populated heading/text columns.")
        if self.layout == "metric" and not self.metric_value and not self.bullets:
            raise ValueError("A metric slide needs its supplied value and label.")
        if self.layout == "table" and self.table is None:
            raise ValueError("A table slide needs actual table data.")
        if self.table and any(len(row) != len(self.table.columns) for row in self.table.rows):
            raise ValueError("Slide table rows must match their columns.")
        return self


class Deck(BaseModel):
    title: str = Field(max_length=150)
    slides: list[Slide] = Field(min_length=1, max_length=15)
    theme: Theme = Field(default_factory=Theme)
    subtitle: str = Field(default="", max_length=200)
    aspect_ratio: Literal["16:9", "4:3"] = "16:9"


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def prepare_sheet(value, request=""):
    sheet = Sheet.model_validate(value)
    numeric = {column for chart in sheet.charts for column in chart.value_columns}
    formats = list(sheet.formats)
    formats.extend("text" for _ in range(len(sheet.columns) - len(sheet.formats)))
    if re.search(r"\b(?:currency|dollars?|usd)\b|\$", request, re.I):
        for column in numeric:
            if column <= len(formats) and formats[column - 1] == "text":
                formats[column - 1] = "eur" if re.search(r"\b(?:euros?|eur)\b|€", request, re.I) else "usd"
    numeric |= {i + 1 for i, fmt in enumerate(formats) if fmt in ("usd", "eur", "decimal", "integer", "percent")}
    for row in sheet.rows:
        for column in numeric:
            if column > len(row):
                continue
            cell = row[column - 1]
            if isinstance(cell, str) and re.fullmatch(r"\s*[$€£]?\s*-?\d+(?:,\d{3})*(?:\.\d+)?%?\s*", cell):
                amount = float(re.sub(r"[^0-9.\-]", "", cell))
                row[column - 1] = amount / 100 if cell.strip().endswith("%") else amount
    for chart in sheet.charts:
        for column in chart.value_columns:
            if not any(column <= len(row) and isinstance(row[column - 1], (int, float)) and not isinstance(row[column - 1], bool) for row in sheet.rows):
                raise ValueError("A chart needs numeric values in its requested data columns.")
    sheet.formats = formats
    return sheet.model_dump()


def write_sheet(value, path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.data_source import AxDataSource, StrRef
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
    sheet = Sheet.model_validate(prepare_sheet(value))
    book = Workbook()
    page = cast(Worksheet, book.active)
    page.title = re.sub(r"[\\/*?:\[\]]", "", sheet.title)[:31] or "MacBot"
    book.properties.title, book.properties.creator = sheet.title, "MacBot"
    page.append([safe_cell(column) for column in sheet.columns])
    for row in sheet.rows:
        if len(row) != len(sheet.columns):
            raise ValueError("Every spreadsheet row must match its column count.")
        page.append([safe_cell(cell) for cell in row])
    for number, column in enumerate(page.columns, start=1):
        width = sheet.column_widths[number - 1] if sheet.column_widths else min(48, max(14, max(len(str(c.value or "")) for c in column) + 2))
        page.column_dimensions[get_column_letter(number)].width = width
    header_fill = sheet.theme.accent if sheet.header_style == "accent" else "FFFFFF" if sheet.header_style == "minimal" else sheet.theme.background
    for cell in page[1]:
        cell.font = Font(name=sheet.theme.heading_face, size=sheet.theme.body_size + 1, bold=True,
                         color=readable_color(sheet.theme.foreground, header_fill))
        cell.fill = PatternFill("solid", fgColor=header_fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=sheet.theme.accent))
    formats = {"integer": "#,##0", "decimal": "#,##0.00", "percent": "0.0%", "usd": '"$"#,##0.00', "eur": '"€"#,##0.00', "date": "yyyy-mm-dd", "text": "@"}
    for row_number, row in enumerate(page.iter_rows(min_row=2), start=2):
        fill = sheet.theme.background if sheet.stripe_rows and row_number % 2 == 0 else "FFFFFF"
        for column_number, cell in enumerate(row, start=1):
            cell.font = Font(name=sheet.theme.font, size=sheet.theme.body_size,
                             color=readable_color(sheet.theme.foreground, fill))
            alignment = sheet.alignments[column_number - 1] if sheet.alignments else "auto"
            if alignment == "auto":
                alignment = "right" if isinstance(cell.value, (int, float)) else "left"
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal=alignment)
            cell.fill = PatternFill("solid", fgColor=fill)
            if column_number <= len(sheet.formats):
                cell.number_format = formats[sheet.formats[column_number - 1]]
    import textwrap
    padding = {"compact": 6, "comfortable": 10, "airy": 16}[sheet.theme.spacing]
    for row in page:
        lines = max(sum(max(1, len(textwrap.wrap(line, width=max(4, int(
            page.column_dimensions[cell.column_letter].width * 11 / sheet.theme.body_size - 2)))))
            for line in str(cell.value if cell.value is not None else "").split("\n")) for cell in row)
        page.row_dimensions[row[0].row].height = min(409, max(24, lines * sheet.theme.body_size * 1.35 + padding))
    page.freeze_panes = "A2"
    if sheet.rows:
        table = Table(displayName="MacBotData", ref=page.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False,
                                              showFirstColumn=False, showLastColumn=False, showColumnStripes=False)
        page.add_table(table)
    else:
        page.auto_filter.ref = page.dimensions
    chart_page = book.create_sheet("Charts") if sheet.charts and sheet.chart_placement == "dashboard" else page
    if chart_page is not page:
        chart_page["A1"] = safe_cell(sheet.title)
        chart_page["A1"].font = Font(name=sheet.theme.heading_face, size=sheet.theme.heading_size,
                                    color=readable_color(sheet.theme.accent, "FFFFFF"), bold=True)
        chart_page.sheet_view.showGridLines = False
    for number, specification in enumerate(sheet.charts):
        if specification.category_column > len(sheet.columns) or any(c < 1 or c > len(sheet.columns) for c in specification.value_columns):
            raise ValueError("Chart columns must refer to existing spreadsheet columns.")
        chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[specification.kind]()
        chart.title = specification.title or sheet.title
        chart.title.overlay = False
        chart.style = 10
        chart_font = CharacterProperties(latin=DrawingFont(typeface=sheet.theme.font),
            sz=int(min(12, sheet.theme.body_size) * 100), solidFill=readable_color(sheet.theme.foreground, "FFFFFF"))
        def chart_text():
            return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=chart_font), endParaRPr=chart_font)])
        if specification.kind != "pie":
            chart.x_axis.axPos, chart.y_axis.axPos = "b", "l"
            for axis in (chart.x_axis, chart.y_axis):
                axis.delete, axis.tickLblPos, axis.txPr = False, "nextTo", chart_text()
            value_format = sheet.formats[specification.value_columns[0] - 1]
            chart.y_axis.numFmt = "General" if value_format == "text" else formats[value_format]
        if chart.legend is not None:
            chart.legend.txPr = chart_text()
            chart.legend.position, chart.legend.overlay = "b", False
        for column in specification.value_columns[:1] if specification.kind == "pie" else specification.value_columns:
            chart.add_data(Reference(page, min_col=column, min_row=1, max_row=len(sheet.rows) + 1), titles_from_data=True)
        categories = Reference(page, min_col=specification.category_column, min_row=2, max_row=len(sheet.rows) + 1)
        chart.set_categories(categories)
        if any(isinstance(row[specification.category_column - 1], str) for row in sheet.rows):
            for series in chart.series:
                series.cat = AxDataSource(strRef=StrRef(f=str(categories)))
        palette = sheet.chart_colors or [sheet.theme.accent, sheet.theme.foreground]
        for index, series in enumerate(chart.series):
            color = palette[index % len(palette)]
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
        if specification.kind == "pie":
            from openpyxl.chart.series import DataPoint
            for index in range(len(sheet.rows)):
                point = DataPoint(idx=index)
                point.graphicalProperties.solidFill = palette[index % len(palette)]
                chart.series[0].data_points.append(point)
        chart.width, chart.height = 19, 10
        anchor = f"{get_column_letter(len(sheet.columns) + 2)}{2 + number * 21}" if sheet.chart_placement == "beside" else f"A{len(sheet.rows) + 4 + number * 22}" if sheet.chart_placement == "below" else f"A{4 + number * 22}"
        chart_page.add_chart(chart, anchor=anchor)
    if sheet.notes:
        notes = book.create_sheet("Notes")
        notes["A1"], notes["A2"] = "About this workbook", safe_cell(sheet.notes)
        notes["A1"].font = Font(name=sheet.theme.heading_face, size=sheet.theme.heading_size, bold=True,
                                color=readable_color(sheet.theme.accent, "FFFFFF"))
        notes["A2"].font = Font(name=sheet.theme.font, size=sheet.theme.body_size)
        notes["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        notes.column_dimensions["A"].width = 85
        notes.row_dimensions[2].height = 120
    page.sheet_view.showGridLines = False
    assert page.sheet_properties.pageSetUpPr is not None
    page.sheet_properties.pageSetUpPr.fitToPage = True
    page.page_setup.orientation, page.page_setup.fitToWidth, page.page_setup.fitToHeight = sheet.print_orientation, 1, 0
    page.print_title_rows = "1:1"
    header, footer = page.oddHeader, page.oddFooter
    assert header is not None and footer is not None
    assert header.center is not None and footer.center is not None
    header.center.text = sheet.title.replace("&", "&&")
    footer.center.text = "Page &P of &N"
    book.save(path)
    checked = load_workbook(path, read_only=True)
    assert checked.active is not None and checked.active.max_row == len(sheet.rows) + 1
    checked.close()


def write_deck(value, path):
    from .slides import render_deck
    render_deck(value, path)


def synthesize(text, path, directory):
    from piper import PiperVoice
    from piper.config import SynthesisConfig
    model = prepare_asset(directory, "voice") / VOICE_PREFIX / VOICE_FILE
    voice = PiperVoice.load(model, use_cuda=False)
    with wave.open(str(path), "wb") as output:
        voice.synthesize_wav(text[:6000], output, syn_config=SynthesisConfig(length_scale=1.05))
    with wave.open(str(path), "rb") as check:
        if check.getnframes() == 0:
            raise ValueError("The voice engine produced an empty recording.")


def draw_image(prompt, path, directory, seed=42, options=None, stop=None):
    import torch
    from diffusers.pipelines.stable_diffusion.pipeline_stable_diffusion import StableDiffusionPipeline
    from diffusers.pipelines.stable_diffusion.pipeline_output import StableDiffusionPipelineOutput
    from PIL import Image, ImageEnhance, ImageOps
    specification = ImageSpec.model_validate(options or {"prompt": prompt, "width": 384, "height": 384})
    scale = min(1, 512 / max(specification.width, specification.height))
    width = max(128, round(specification.width * scale / 8) * 8)
    height = max(128, round(specification.height * scale / 8) * 8)
    torch.set_num_threads(4)
    model = prepare_asset(directory, "image")
    pipeline = StableDiffusionPipeline.from_pretrained(str(model), torch_dtype=torch.float32,
                                                      local_files_only=True)
    pipeline.to("cpu")
    pipeline.enable_attention_slicing()
    pipeline.set_progress_bar_config(disable=True)
    try:
        styles = {"photographic": "natural photography, realistic lighting", "anime": "anime illustration, expressive linework",
                  "watercolor": "watercolor on textured paper", "cinematic": "cinematic lighting, careful composition",
                  "pixel-art": "pixel art, crisp limited palette", "vintage": "vintage film photograph, warm tones",
                  "illustration": "editorial illustration, deliberate composition", "none": ""}
        def check_stop(pipeline, step, timestep, values):
            if stop is not None and stop.is_set():
                raise InterruptedError("Image creation stopped.")
            return values
        # Diffusers annotates this as a legacy callback, but step-end hooks receive four arguments.
        result = pipeline(prompt=(specification.prompt + ", " + styles[specification.style])[:1000], negative_prompt=specification.negative_prompt,
                          width=width, height=height, num_inference_steps=18, guidance_scale=7.0,
                          generator=torch.Generator(device="cpu").manual_seed(seed), callback_on_step_end=cast(Any, check_stop))
        assert isinstance(result, StableDiffusionPipelineOutput)
        if result.nsfw_content_detected and result.nsfw_content_detected[0]:
            raise ValueError("The image model's safety check declined this output. Try another description.")
        picture = result.images[0]
        assert isinstance(picture, Image.Image)
        if specification.filter == "grayscale":
            picture = ImageOps.grayscale(picture).convert("RGB")
        elif specification.filter == "sepia":
            picture = ImageOps.colorize(ImageOps.grayscale(picture), "#24140c", "#f8e3b7")
        elif specification.filter == "vivid":
            picture = ImageEnhance.Color(picture).enhance(1.35)
        if stop is not None and stop.is_set():
            raise InterruptedError("Image creation stopped.")
        picture = ImageOps.fit(picture, (specification.width, specification.height), method=Image.Resampling.LANCZOS)
        picture.save(path, format=specification.format.upper(), **({"quality": 95} if specification.format != "png" else {}))
        with Image.open(path) as check:
            check.verify()
        return {**specification.model_dump(), "native_width": width, "native_height": height, "seed": seed,
                "upscaled": width != specification.width or height != specification.height}
    finally:
        del pipeline
        gc.collect()


async def generate(mode, prompt, runtime, model, store, job_id, emit, uploads=None):
    import asyncio
    metadata = {}
    extension = {"image": "png", "audio": "wav", "spreadsheet": "xlsx", "presentation": "pptx", "document": "docx"}[mode]
    artifact_id = uuid5(NAMESPACE_URL, f"macbot:{job_id}:{mode}").hex
    path = store.directory / "artifacts" / f"{artifact_id}.{extension}"
    emit({"type": "stage", "agent": "creator", "label": f"Creating your {mode}"})
    if mode == "image":
        raw = await runtime.complete(model, "Translate this request into a concrete image brief. Preserve explicit size, aspect ratio, style, filter and file format. "
            "Enhance composition and lighting without changing the subject. Width and height are export pixels, 128 to 2048; native rendering is limited to 512 on the longest edge. "
            "Return actual values, not a schema. Request:\n" + prompt, json_mode=ImageSpec.model_json_schema(), max_tokens=550)
        specification = ImageSpec.model_validate_json(raw)
        request = prompt.split("<untrusted_attachment", 1)[0]
        dimensions = re.search(r"\b(\d{2,4})\s*[x×]\s*(\d{2,4})\b", request, re.I)
        if dimensions:
            specification.width, specification.height = map(int, dimensions.groups())
            specification = ImageSpec.model_validate(specification.model_dump())
        ratio = re.search(r"\b(16:9|9:16|4:3|3:4|1:1)\b", request)
        if ratio and not dimensions:
            a, b = map(int, ratio[1].split(":"))
            specification.width, specification.height = (1024, round(1024 * b / a)) if a >= b else (round(1024 * a / b), 1024)
        for file_format in ("webp", "jpeg", "png"):
            if re.search(r"\b" + file_format + r"\b", request, re.I):
                specification.format = cast(Literal["png", "jpeg", "webp"], file_format)
        await runtime.unload(model)
        extension = specification.format
        path = path.with_suffix("." + extension)
        import threading
        stop = threading.Event()
        worker = asyncio.create_task(asyncio.to_thread(draw_image, specification.prompt, path, store.directory,
            options=specification.model_dump(), stop=stop))
        try:
            metadata = await asyncio.shield(worker)
        except asyncio.CancelledError:
            stop.set()
            try:
                await worker
            except InterruptedError:
                pass
            raise
        answer = f"Your {extension.upper()} image is ready: **{specification.width} × {specification.height}**."
        if metadata["upscaled"]:
            answer += f" Rendered locally at {metadata['native_width']} × {metadata['native_height']} and resized to your export size."
    elif mode == "audio":
        single = bool(re.search(r"\b(?:one|single)(?: short)? sentence\b", prompt, re.I))
        limit = "exactly one sentence of at most 25 words" if single else "at most 200 words"
        text = await runtime.complete(model, "Write an English voice script for this request. Return only the spoken words, no Markdown, " + limit + ":\n" + prompt, max_tokens=120 if single else 400)
        if single:
            text = re.split(r"(?<=[.!?])\s+", text.strip(), maxsplit=1)[0]
            text = " ".join(text.split()[:25])
        else:
            text = " ".join(text.split()[:200])
        if not text.strip():
            raise ValueError("The model returned an empty voice script. Try another request.")
        await runtime.unload(model)
        await asyncio.to_thread(synthesize, text, path, store.directory)
        answer = text
    else:
        schema = Sheet if mode == "spreadsheet" else Deck if mode == "presentation" else DocumentSpec
        example = {"title": "Budget", "columns": ["Item", "Cost"], "rows": [["Coffee", 5]], "formats": ["text", "usd"],
            "charts": [{"kind": "bar", "title": "Costs", "category_column": 1, "value_columns": [2]}]} if mode == "spreadsheet" else {
            "title": "Workshop",
            "slides": [{"title": "Introduction", "layout": "cover", "subtitle": "Music together", "bullets": [], "notes": "Introduce the topic."},
                       {"title": "Compare options", "layout": "comparison", "columns": [{"heading": "Option A", "text": "First approach"}, {"heading": "Option B", "text": "Second approach"}], "bullets": []},
                       {"title": "A supplied figure", "layout": "metric", "metric_value": "30", "metric_label": "Participants", "bullets": []}]}
        if mode == "document":
            example = {"title": "Workshop guide", "format": "both", "sections": [{"heading": "Getting started", "text": "Bring your instrument and a notebook.", "bullets": []}]}
        instruction = "Create polished, useful file content and a coherent design based on the user's request. Honour colours, tone, structure, format and page preferences. "
        instruction += ("Choose the visual direction for this particular audience and content before producing the JSON. "
            "Select editorial, minimal, formal or bold style; an intentional palette; body and heading fonts; "
            "body, heading and title point sizes; compact, comfortable or airy spacing; and line spacing. "
            "Use a coherent font pairing when appropriate, and a single font when requested. Do not reuse a default palette for every request. "
            "Available fonts: " + ", ".join(FONT_FILES) + ". Respect explicit user fonts, sizes and colours. ")
        from datetime import date
        instruction += f"Today's local date is {date.today().isoformat()}. Do not invent dates, recipients, named organisations, results or statistics. Omit unspecified dates. "
        instruction += "Use facts provided by the user; otherwise clearly label estimates. Return actual data values, never a JSON schema. Use six-digit hex theme colours without #. "
        if mode == "spreadsheet":
            instruction += "Set number formats for each column. Include charts only when requested or clearly useful. Keep headers concise. "
            instruction += "Use body_size 10-13. Choose header_style, row stripes, column widths and alignments suitable for the data. Put charts beside, below or on a dashboard. Choose distinct chart_colors from a palette that complements the workbook theme. "
        elif mode == "presentation":
            instruction += "Use varied cover, comparison, timeline, metric and split layouts where content fits. Avoid repeated bullet slides. "
            instruction += "Write short bullets; put detailed explanations in speaker notes. Columns contain heading and text. Never invent quotations or numerical results. "
            instruction += "Fill comparison columns and timeline steps with useful content, not empty arrays. Suggested teaching activities are welcome when clearly presented as a proposal. "
            instruction += "Use body_size 18-24, heading_size 28-36 and title_size 40-56. Alternate base, accent or contrast slide backgrounds with purpose. Use cards, section dividers and native table layouts when helpful. Tables have at most 5 columns and 6 rows. Keep slides concise so they fit; preserve details in notes. "
        else:
            instruction += "Organise the document into readable sections. Put lists in bullets and tables in table.columns and table.rows, never as Markdown in text. "
            instruction += "The renderer creates the title page from title and subtitle when cover=true; do not add a Title Page or Cover section. Use a cover unless a short document was requested. "
            instruction += "Use body_size 10-12, heading_size 16-24 and title_size 28-40. Choose paper, orientation and margin. Use callout emphasis for key sections and page breaks only where the content needs them. "
        instruction += "Example content object: " + json.dumps(example)
        output_schema = schema.model_json_schema()
        output_schema["required"] = list(dict.fromkeys([*output_schema.get("required", []), "theme"]))
        output_schema["$defs"]["Theme"]["required"] = ["background", "foreground", "accent", "font", "heading_font",
            "style", "body_size", "heading_size", "title_size", "line_spacing", "spacing"]
        value: dict[str, Any] | None = None
        for attempt in range(2):
            raw = await runtime.complete(model, instruction + "\nRequest:\n" + prompt, json_mode=output_schema, max_tokens=4200)
            try:
                value = schema.model_validate_json(raw).model_dump()
                break
            except ValueError as error:
                if attempt == 1:
                    raise ValueError("The model could not create a usable file structure. Try a shorter, more specific request.") from error
                emit({"type": "stage", "agent": "creator", "label": "Improving the file structure"})
                instruction += "\nYour previous output was invalid. Include actual content values for every required field. Validation feedback: " + str(error)[:600]
        assert value is not None
        request = prompt.split("<untrusted_attachment", 1)[0]
        value["theme"] = requested_theme(value["theme"], request)
        if mode == "spreadsheet":
            value = prepare_sheet(value, request)
        elif mode == "presentation":
            layouts = re.findall(r"\b(cover|comparison|timeline|metric|quote|split)\b", request, re.I)
            if len(layouts) == len(value["slides"]):
                for slide, layout in zip(value["slides"], layouts):
                    slide["layout"] = layout.lower()
            metric = re.search(r"\b(\d[\d,.]*)\s+(participants|attendees|people|students|members)\b", request, re.I)
            if metric:
                for slide in value["slides"]:
                    if slide["layout"] == "metric":
                        slide["metric_value"], slide["metric_label"] = metric.groups()
        if mode == "document":
            request = prompt.split("<untrusted_attachment", 1)[0]
            from .design import prepare_document
            value = prepare_document(value, request)
            if re.search(r"\bpdf\b", request, re.I):
                value["format"] = "both" if re.search(r"\b(?:docx|word)\b", request, re.I) else "pdf"
            elif re.search(r"\b(?:docx|word)\b", request, re.I):
                value["format"] = "docx"
            if re.search(r"\blandscape\b", request, re.I):
                value["orientation"] = "landscape"
        if mode == "document":
            formats = ["docx", "pdf"] if value["format"] == "both" else [value["format"]]
            artifacts = []
            for fmt in formats:
                target = path.with_suffix("." + fmt)
                await asyncio.to_thread(write_document, value, target, fmt)
                artifacts.append({"id": uuid5(NAMESPACE_URL, f"macbot:{job_id}:document:{fmt}").hex,
                    "kind": "document", "name": f"MacBot.{fmt}", "path": str(target), "metadata": {"theme": value["theme"], "title": value["title"]}})
            return {"answer": f"Created **{value['title']}** with your document layout.", "sources": [], "artifacts": artifacts}
        if mode == "presentation":
            from .slides import render_deck
            images = {}
            for identifier in uploads or []:
                rows = store.execute("SELECT kind,path FROM uploads WHERE id=?", (identifier,))
                if rows and rows[0]["kind"] == "image":
                    images[identifier] = rows[0]["path"]
            await asyncio.to_thread(render_deck, value, path, images)
        else:
            await asyncio.to_thread(write_sheet, value, path)
        metadata = {"theme": value["theme"]}
        answer = f"Created **{value['title']}**. Download your {mode} below."
    return {"answer": answer, "sources": [], "artifacts": [{"id": artifact_id, "kind": mode,
            "name": f"MacBot.{extension}", "path": str(path), "metadata": {"prompt": prompt, "engine": "tiny-sd" if mode == "image" else "Piper" if mode == "audio" else model,
            **metadata}}]}
