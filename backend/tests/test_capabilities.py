import json

import pytest
from langgraph.checkpoint.sqlite.aio import AsyncSqliteSaver
from langgraph.types import Command
from langchain_core.runnables import RunnableConfig

from macbot.artifacts import write_deck, write_sheet
from macbot.evidence import render_report
from macbot.graph import State, build_graph
from macbot.mcp_client import ServerConfig, add_workspace, create_server, execute_tool, fingerprint, inspect_server
from macbot.secrets import open_config
from macbot.store import Store, now


def test_generated_office_files_are_readable_and_formula_safe(tmp_path):
    from openpyxl import load_workbook
    from openpyxl.worksheet.worksheet import Worksheet
    from pptx import Presentation
    sheet = tmp_path / "sheet.xlsx"
    write_sheet({"title": "Budget", "columns": ["Item", "Cost"], "rows": [["=HYPERLINK(\"bad\")", 17]]}, sheet)
    book = load_workbook(sheet)
    page = book.active
    assert isinstance(page, Worksheet)
    assert str(page["A2"].value).startswith("'")
    assert page["B2"].value == 17
    book.close()
    deck = tmp_path / "deck.pptx"
    write_deck({"title": "Plan", "slides": [{"title": "Next steps", "bullets": ["Build", "Review"]}]}, deck)
    assert len(Presentation(deck).slides) == 1
    with pytest.raises(ValueError):
        write_sheet({"title": "Invalid", "columns": ["One"], "rows": [[1, 2]]}, sheet)


def test_job_completion_is_atomic_and_repeatable(tmp_path):
    store = Store(tmp_path)
    chat = store.create_chat()["id"]
    store.execute("INSERT INTO jobs(id,chat_id,kind,status,request,created) VALUES(?,?,?,?,?,?)", ("same", chat, "chat", "running", "{}", now()))
    first = store.finish_job("same", chat, {"answer": "Ready", "sources": []})
    assert store.finish_job("same", chat, {"answer": "Ready", "sources": []}) == first
    assert len(store.messages(chat)) == 1
    assert store.job("same")["status"] == "completed"
    store.db.close()


def test_mcp_credentials_are_encrypted_and_redacted(tmp_path):
    store = Store(tmp_path)
    public = create_server(store, ServerConfig(name="Private", command="server.exe", env={"KEY": "secret-123"}))
    saved = json.loads(store.execute("SELECT data FROM mcp_servers WHERE id=?", (public["id"],))[0]["data"])
    assert "secret-123" not in json.dumps(saved)
    assert "secret-123" not in json.dumps(public)
    assert open_config(saved)["env"]["KEY"] == "secret-123"
    store.db.close()


async def test_real_workspace_mcp_is_scoped_and_replay_safe(tmp_path):
    store = Store(tmp_path)
    add_workspace(store)
    config = ServerConfig(name="Workspace", transport="workspace", enabled=True).model_dump()
    store.execute("UPDATE mcp_servers SET data=? WHERE id='workspace'", (json.dumps(config),))
    tools = await inspect_server(config, tmp_path)
    assert {tool["name"] for tool in tools} >= {"list_files", "read_file", "write_file"}
    proposal = {"server_id": "workspace", "config_hash": fingerprint(config), "tool": "write_file", "arguments": {"name": "approved.txt", "content": "approved content"}}
    result = await execute_tool(store, "job1", proposal)
    assert (tmp_path / "workspace/approved.txt").read_text() == "approved content"
    assert await execute_tool(store, "job1", proposal) == result
    escape = {**proposal, "arguments": {"name": "../outside.txt", "content": "blocked"}}
    with pytest.raises(ValueError):
        await execute_tool(store, "job2", escape)
    assert not (tmp_path / "outside.txt").exists()
    store.db.close()


async def test_graph_approval_persists_before_any_tool_effect(tmp_path):
    store = Store(tmp_path)
    add_workspace(store)
    config = ServerConfig(name="Workspace", transport="workspace", enabled=True).model_dump()
    store.execute("UPDATE mcp_servers SET data=? WHERE id='workspace'", (json.dumps(config),))
    class Runtime:
        async def complete(self, *args, **kwargs):
            return json.dumps({"server_id": "workspace", "tool": "write_file", "arguments": {"name": "gated.txt", "content": "After approval"}})
    chat = store.create_chat()["id"]
    state: State = {"prompt": "Create a file", "mode": "tools", "chat_id": chat, "model": "local", "job_id": "gated"}
    run_config: RunnableConfig = {"configurable": {"thread_id": "gated"}}
    async with AsyncSqliteSaver.from_conn_string(str(tmp_path / "checkpoints.db")) as saver:
        graph = build_graph(Runtime(), store, lambda _: None, checkpointer=saver)
        paused = await graph.ainvoke(state, run_config)
        assert paused.get("__interrupt__")
        assert not (tmp_path / "workspace/gated.txt").exists()
    async with AsyncSqliteSaver.from_conn_string(str(tmp_path / "checkpoints.db")) as saver:
        graph = build_graph(Runtime(), store, lambda _: None, checkpointer=saver)
        result = await graph.ainvoke(Command(resume=True), run_config)
        assert "completed" in result["answer"]
    assert (tmp_path / "workspace/gated.txt").exists()
    store.db.close()


def test_report_abstains_without_supported_claims():
    assert "could not establish" in render_report([], 2)


def test_unicode_pdf_and_presentation_ingestion(tmp_path):
    from macbot.documents import extract, export_document
    path = tmp_path / "unicode.pdf"
    export_document("Résumé for José: 37 items.", path, "pdf")
    assert "José" in extract(path.name, path.read_bytes())[1]
    deck = tmp_path / "slides.pptx"
    write_deck({"title": "Plan", "slides": [{"title": "Workshop", "bullets": ["Bring 17 records"]}]}, deck)
    assert "Bring 17 records" in extract(deck.name, deck.read_bytes())[1]
