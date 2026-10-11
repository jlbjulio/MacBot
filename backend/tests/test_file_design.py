from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from typing import Any, cast

from macbot.artifacts import prepare_sheet, write_sheet
from macbot.artifacts import Deck
import pytest


def test_currency_strings_become_real_chart_values(tmp_path):
    value = {"title": "Budget", "columns": ["Item", "Cost"],
             "rows": [["Coffee", "$5.00"], ["Books", "$12.00"]],
             "charts": [{"value_columns": [2]}]}
    output = tmp_path / "budget.xlsx"
    write_sheet(prepare_sheet(value, "Use currency formatting and a bar chart"), output)
    book = load_workbook(output)
    page = cast(Worksheet, book.active)
    assert page["B2"].value == 5
    assert page["B3"].value == 12
    assert "$" in page["B2"].number_format
    assert len(cast(Any, page)._charts) == 1
    book.close()


def test_empty_content_slides_are_rejected():
    with pytest.raises(ValueError, match="meaningful"):
        Deck.model_validate({"title": "Workshop", "slides": [{"title": "Comparison", "layout": "comparison", "bullets": []}]})


def test_markdown_model_body_becomes_formatted_document(tmp_path):
    from docx import Document
    from pypdf import PdfReader
    from macbot.design import write_document

    value = {"title": "Workshop", "sections": [
        {"heading": "Title Page", "text": "Prepared for the community"},
        {"heading": "Objectives", "text": "* **Skills:** Learn together.\n* Share music."},
        {"heading": "Timeline", "text": "| Phase | Activity |\n| :--- | ---: |\n| 1 | **Planning** |\n| 2 | Workshop |"},
    ]}
    word, pdf = tmp_path / "workshop.docx", tmp_path / "workshop.pdf"
    write_document(value, word, "docx")
    write_document(value, pdf, "pdf")
    doc = Document(word)
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 1).text == "Planning"
    assert any(p.style is not None and p.style.name == "List Bullet" and p.text.startswith("Skills:")
               and p.runs[0].bold for p in doc.paragraphs)
    assert "Title Page" not in [p.text for p in doc.paragraphs]
    assert any(p.text == "Prepared for the community" for p in doc.paragraphs)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    assert "Planning" in text and "**" not in text and "|" not in text


def test_requested_timeline_is_a_native_table_when_model_returns_bullets(tmp_path):
    from docx import Document
    from docx.styles.style import ParagraphStyle
    from macbot.design import prepare_document, requested_theme, write_document

    plan = {"title": "Workshop", "sections": [
        {"heading": "Objectives", "text": "Learn together."},
        {"heading": "Implementation Timeline", "text": "Project phases.",
         "bullets": ["Month 1: Recruitment", "Month 2: Workshop"]},
    ]}
    formatted = prepare_document(plan, "Create a proposal with a small timeline table.")
    formatted["theme"] = requested_theme(formatted["theme"], "Use a purple and cream theme with Georgia.")
    path = tmp_path / "workshop.docx"
    write_document(formatted, path, "docx")
    document = Document(path)
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "Month 1"
    assert document.tables[0].cell(2, 1).text == "Workshop"
    assert cast(ParagraphStyle, document.styles["Normal"]).font.name == "Georgia"
    assert "674183" in document.tables[0]._tbl.xml
    assert "Month 1: Recruitment" not in [p.text for p in document.paragraphs]
    assert prepare_document(plan, "Create a proposal without tables.")["sections"][1]["table"] is None


def test_requested_typography_survives_word_and_pdf_export(tmp_path):
    from docx import Document
    from docx.styles.style import ParagraphStyle
    from pypdf import PdfReader
    from macbot.design import requested_theme, write_document

    value = {"title": "Community arts", "cover": False, "margin": 0.7,
             "theme": requested_theme({}, "Use headings in Georgia, body in Arial, body size 12pt and title size 36pt."),
             "sections": [{"heading": "Plan", "text": "A guide for the community.", "emphasis": "callout"},
                          {"heading": "Delivery", "text": "All activities remain visible.", "page_break_before": True}]}
    word, pdf = tmp_path / "arts.docx", tmp_path / "arts.pdf"
    write_document(value, word, "docx")
    write_document(value, pdf, "pdf")
    document = Document(word)
    assert cast(ParagraphStyle, document.styles["Normal"]).font.name == "Arial"
    assert cast(ParagraphStyle, document.styles["Normal"]).font.size.pt == 12
    assert cast(ParagraphStyle, document.styles["Title"]).font.name == "Georgia"
    assert "asciiTheme" not in cast(ParagraphStyle, document.styles["Title"]).element.xml
    assert cast(ParagraphStyle, document.styles["Title"]).font.size.pt == 36
    assert document.sections[0].left_margin.inches == pytest.approx(0.7)
    pages = PdfReader(pdf).pages
    assert len(pages) == 2
    assert "All activities remain visible." in pages[1].extract_text()
    fonts = {str(font.get_object()["/BaseFont"]) for page in pages
             for font in page["/Resources"]["/Font"].get_object().values()}
    assert any("Georgia" in name for name in fonts)
    assert any("Arial" in name for name in fonts)


def test_workbook_design_keeps_wrapped_content_and_real_dashboard_data(tmp_path):
    value = {"title": "Arts budget", "columns": ["Activity", "Budget"],
             "rows": [["A long community workshop description that needs multiple lines in the requested narrow column", 125],
                      ['=HYPERLINK("https://example.com")', 250]],
             "theme": {"font": "Arial", "heading_font": "Georgia", "body_size": 12, "spacing": "airy"},
             "column_widths": [24, 16], "alignments": ["left", "right"],
             "header_style": "accent", "chart_placement": "dashboard", "chart_colors": ["674183", "236B70"],
             "charts": [{"kind": "bar", "value_columns": [2]}], "formats": ["text", "usd"]}
    path = tmp_path / "arts.xlsx"
    write_sheet(value, path)
    book = load_workbook(path)
    page = cast(Worksheet, book.active)
    assert page["A2"].value == value["rows"][0][0]
    assert page["A3"].data_type == "s" and str(page["A3"].value).startswith("'=")
    assert page.row_dimensions[2].height > 50
    assert page["A1"].font.name == "Georgia" and page["A2"].font.name == "Arial"
    assert page.column_dimensions["A"].width == 24
    assert page.tables["MacBotData"].tableStyleInfo.name == "TableStyleLight1"
    assert page.auto_filter.ref is None
    assert page.tables["MacBotData"].autoFilter.ref == page.dimensions
    chart = cast(Any, book["Charts"])._charts[0]
    assert "Arts budget" in chart.series[0].val.numRef.f
    assert chart.series[0].graphicalProperties.solidFill.srgbClr == "674183"
    assert chart.x_axis.delete is False and chart.x_axis.tickLblPos == "nextTo"
    assert chart.series[0].cat.strRef.f.endswith("$A$2:$A$3")
    book.close()


def test_slides_keep_all_eight_items_and_editable_tables_within_page(tmp_path):
    from pptx import Presentation
    from macbot.slides import render_deck

    items = [f"Activity {number}: a short useful description." for number in range(1, 9)]
    value = {"title": "Arts workshop", "aspect_ratio": "4:3",
             "theme": {"font": "Arial", "heading_font": "Georgia", "body_size": 20, "heading_size": 30,
                       "title_size": 44, "style": "bold", "accent": "674183"},
             "slides": [{"title": "Activities", "layout": "cards", "bullets": items},
                        {"title": "Schedule", "layout": "table", "background": "contrast", "bullets": [],
                         "table": {"columns": ["Phase", "Activity"], "rows": [["1", "Planning"], ["2", "Workshop"]]}}]}
    path = tmp_path / "arts.pptx"
    render_deck(value, path)
    presentation = Presentation(path)
    slide_text = "\n".join(shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame)
    assert all(item in slide_text for item in items)
    assert presentation.slides[1].background.fill.fore_color.rgb != presentation.slides[0].background.fill.fore_color.rgb
    tables = [shape.table for shape in presentation.slides[1].shapes if shape.has_table]
    assert len(tables) == 1 and tables[0].cell(2, 1).text == "Workshop"
    assert any(shape.has_text_frame and any(p.font.name == "Georgia" for p in shape.text_frame.paragraphs)
               for shape in presentation.slides[0].shapes)
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert 0 <= shape.left and shape.left + shape.width <= presentation.slide_width
            assert 0 <= shape.top and shape.top + shape.height <= presentation.slide_height


@pytest.mark.parametrize("theme", [{"font": "../../payload"}, {"accent": '<script>alert(1)</script>'}, {"body_size": 10000}])
def test_file_design_rejects_unsafe_or_unbounded_style_values(theme):
    from macbot.design import Theme

    with pytest.raises(ValueError):
        Theme.model_validate(theme)
