import io
from types import SimpleNamespace

from PIL import Image
import pytest
from tokenizers import Tokenizer
from tokenizers.models import WordLevel
from tokenizers.pre_tokenizers import Whitespace

from macbot.documents import extract
from macbot.retrieval import RetrievalIndex


def test_chunks_keep_pages_slides_and_timestamps_separate(tmp_path):
    index = RetrievalIndex(tmp_path)
    index.tokenizer = Tokenizer(WordLevel({"[UNK]": 0}, unk_token="[UNK]"))
    index.tokenizer.pre_tokenizer = Whitespace()
    chunks = list(index.chunks("Page 1\nFirst result\n\nPage 2\nSecond result\n\nSlide 3\nThird result\n\nAudio segment 0\nOpening speech"))
    assert [c["page"] for c in chunks[:2]] == [1, 2]
    assert "Second" not in chunks[0]["text"]
    assert chunks[2]["slide"] == 3
    assert chunks[3]["timestamp"] == 0 and chunks[3]["text"].endswith("Opening speech")


def test_duplicate_vectors_are_copied_without_model_inference(tmp_path):
    from qdrant_client import QdrantClient, models
    index = RetrievalIndex(tmp_path)
    index.client = QdrantClient(":memory:")
    index.client.create_collection(index.collection, vectors_config=models.VectorParams(size=2, distance=models.Distance.COSINE))
    index.client.upsert(index.collection, points=[models.PointStruct(id=1, vector=[1.0, 0.0],
        payload={"upload_id": "original", "name": "original.pdf", "text": "Invoice 837.29", "page": 2, "chunk": 0})])
    assert index.copy_upload("original", "copy", "renamed.pdf") == 1
    points, _ = index.client.scroll(index.collection, scroll_filter=models.Filter(must=[
        models.FieldCondition(key="upload_id", match=models.MatchValue(value="copy"))]), with_vectors=True)
    assert points[0].payload["name"] == "renamed.pdf" and points[0].payload["page"] == 2
    assert points[0].payload["text"] == "Invoice 837.29" and points[0].vector == [1.0, 0.0]
    assert index.client.count(index.collection).count == 2
    index.close()


def test_office_images_and_notes_reach_document_text(tmp_path, monkeypatch):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from macbot import ocr
    class Reader:
        def __init__(self, *args):
            pass
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
        def read(self, image):
            return "Scanned invoice total 837.29"
    monkeypatch.setattr(ocr, "DocumentOCR", Reader)
    image = io.BytesIO()
    Image.new("RGB", (200, 100), "white").save(image, format="PNG")
    doc = Document()
    doc.add_paragraph("Native document heading")
    doc.add_picture(io.BytesIO(image.getvalue()))
    path = tmp_path / "invoice.docx"
    doc.save(path)
    text = extract(path.name, path.read_bytes(), tmp_path)[1]
    assert "Native document heading" in text and "837.29" in text
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(image.getvalue()), 0, 0, width=Inches(2))
    slide.notes_slide.notes_text_frame.text = "Speaker notes: payment is due Friday"
    path = tmp_path / "invoice.pptx"
    deck.save(path)
    text = extract(path.name, path.read_bytes(), tmp_path)[1]
    assert "Slide 1" in text and "837.29" in text and "payment is due Friday" in text


def test_audio_segments_keep_their_own_transcript(tmp_path, monkeypatch):
    import numpy as np
    from macbot import media
    index = RetrievalIndex(tmp_path)
    saved = []
    class Embeddings:
        def encode(self, *args, **kwargs):
            return [np.array([1.0, 0.0])]
    index.gemma = Embeddings()
    index.client = SimpleNamespace(upsert=lambda collection, points: saved.extend(points))
    monkeypatch.setattr(media, "audio_waveform", lambda _: np.zeros(60 * 16000, dtype=np.float32))
    segments = [{"start": 0, "end": 20, "text": "Opening remarks"}, {"start": 35, "end": 55, "text": "Closing balance 837.29"}]
    assert index.index_media("audio", "meeting.wav", "audio", tmp_path / "meeting.wav", "whole transcript", segments) == 2
    assert saved[0].payload["text"] == "Opening remarks"
    assert saved[1].payload["text"] == "Closing balance 837.29" and saved[1].payload["timestamp"] == 30


def test_complete_book_keeps_the_final_page_and_retrieval_location(tmp_path):
    from reportlab.pdfgen.canvas import Canvas

    content = io.BytesIO()
    canvas = Canvas(content)
    for page in range(1, 311):
        canvas.setFont("Helvetica", 8)
        for line in range(28):
            canvas.drawString(36, 780 - line * 20, f"Chapter {page}: " + "Readable book content. " * 3)
        if page == 310:
            canvas.drawString(36, 190, "Final appendix total 837.29")
        canvas.showPage()
    canvas.save()
    text = extract("book.pdf", content.getvalue(), tmp_path)[1]
    assert len(text) > 200_000 and text.rstrip().endswith("Final appendix total 837.29")
    index = RetrievalIndex(tmp_path)
    index.tokenizer = Tokenizer(WordLevel({"[UNK]": 0}, unk_token="[UNK]"))
    index.tokenizer.pre_tokenizer = Whitespace()
    final = next(chunk for chunk in index.chunks(text) if "Final appendix total" in chunk["text"])
    assert final["page"] == 310


def test_long_scan_reuses_one_reader_and_keeps_every_page(tmp_path, monkeypatch):
    from reportlab.pdfgen.canvas import Canvas
    from macbot import ocr
    from macbot.media import scan_pdf

    counts = {"sessions": 0, "pages": 0, "closed": 0}
    class Reader:
        def __init__(self, *args):
            counts["sessions"] += 1
        def __enter__(self):
            return self
        def __exit__(self, *args):
            counts["closed"] += 1
        def read(self, image):
            counts["pages"] += 1
            return f"Recognized scan {counts['pages']}"
    monkeypatch.setattr(ocr, "DocumentOCR", Reader)
    content = io.BytesIO()
    canvas = Canvas(content)
    for _ in range(42):
        canvas.rect(36, 36, 10, 10, fill=1)
        canvas.showPage()
    canvas.save()
    text = extract("scanned-book.pdf", content.getvalue(), tmp_path)[1]
    assert "Page 42\n" in text and text.rstrip().endswith("Recognized scan 42")
    assert counts == {"sessions": 1, "pages": 42, "closed": 1}
    counts.update(sessions=0, pages=0, closed=0)
    text = scan_pdf(content.getvalue(), directory=tmp_path)
    assert "Page 42\n" in text and counts == {"sessions": 1, "pages": 42, "closed": 1}


def test_office_keeps_content_beyond_old_limits_in_source_order(tmp_path, monkeypatch):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from macbot import ocr

    class Reader:
        def __init__(self, *args):
            pass
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
        def read(self, image):
            return "Embedded scan marker"
    monkeypatch.setattr(ocr, "DocumentOCR", Reader)
    document = Document()
    document.add_paragraph("Before the table")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table in source order"
    document.add_paragraph("After the table")
    for _ in range(30):
        document.add_paragraph("Full document content. " * 400)
    for color in range(31):
        picture = io.BytesIO()
        Image.new("RGB", (16, 16), (color, 0, 0)).save(picture, format="PNG")
        document.add_picture(io.BytesIO(picture.getvalue()))
    document.add_paragraph("Final document paragraph")
    content = io.BytesIO()
    document.save(content)
    text = extract("complete.docx", content.getvalue(), tmp_path)[1]
    assert len(text) > 200_000 and "Final document paragraph" in text
    assert text.index("Before the table") < text.index("Table in source order") < text.index("After the table")
    assert text.count("Embedded scan marker") == 31
    deck = Presentation()
    for number in range(101):
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = f"Section {number + 1}"
    content = io.BytesIO()
    deck.save(content)
    assert "Slide 101\nSection 101" in extract("complete.pptx", content.getvalue(), tmp_path)[1]
    workbook = Workbook()
    workbook.active.cell(1001, 51, "Last row and column marker")
    content = io.BytesIO()
    workbook.save(content)
    assert "Last row and column marker" in extract("complete.xlsx", content.getvalue(), tmp_path)[1]


def test_extraction_rejects_unsafe_expansion_instead_of_cutting_text(tmp_path, monkeypatch):
    from macbot import documents

    monkeypatch.setattr(documents, "MAX_EXTRACTED_CHARACTERS", 32)
    with pytest.raises(ValueError, match="content has not been truncated"):
        extract("too-large.txt", b"unbounded content " * 10, tmp_path)


def test_index_saves_batches_before_consuming_the_whole_document(tmp_path, monkeypatch):
    import numpy as np

    index = RetrievalIndex(tmp_path)
    consumed, saved = [], []
    def chunks(text):
        for number in range(40):
            consumed.append(number)
            yield {"text": f"Evidence {number}", "page": number + 1, "offset": number}
    def save(collection, points):
        if not saved:
            assert len(consumed) == 16
        saved.extend(points)
    index.client = SimpleNamespace(upsert=save)
    index.gemma = SimpleNamespace(encode=lambda values: [np.array([1.0, 0.0]) for _ in values])
    monkeypatch.setattr(index, "chunks", chunks)
    assert index.index("book", "book.pdf", "") == 40
    assert len({point.id for point in saved}) == 40
    assert saved[-1].payload["page"] == 40 and saved[-1].payload["chunk"] == 39
