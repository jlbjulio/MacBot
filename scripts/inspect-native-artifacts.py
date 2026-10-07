import json
import sqlite3
import wave
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pptx import Presentation
from pptx.shapes.autoshape import Shape

root = Path(__file__).resolve().parents[1]
directory = Path(json.loads((root / "logs/quality/native/session.json").read_text())["profile"])
db = sqlite3.connect(directory / "macbot.sqlite3")
db.row_factory = sqlite3.Row
results = []
for item in db.execute("SELECT * FROM artifacts ORDER BY created"):
    path = Path(item["path"])
    result = {"kind": item["kind"], "name": item["name"], "bytes": path.stat().st_size}
    if item["kind"] == "spreadsheet":
        workbook = load_workbook(path)
        worksheet = cast(Worksheet, workbook.active)
        result["cells"] = list(worksheet.values)
        result["charts"] = len(cast(Any, worksheet)._charts)
        result["formats"] = [cell.number_format for cell in worksheet[2]]
        result["header_colour"] = worksheet["A1"].fill.fgColor.rgb
        workbook.close()
    elif item["kind"] == "presentation":
        presentation = Presentation(str(path))
        width, height = presentation.slide_width, presentation.slide_height
        assert width is not None and height is not None
        result["slides"] = [[shape.text for shape in slide.shapes if isinstance(shape, Shape) and shape.has_text_frame] for slide in presentation.slides]
        result["notes"] = [frame.text if (frame := slide.notes_slide.notes_text_frame) is not None else "" for slide in presentation.slides]
        result["bounds_valid"] = all(shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= width + 10000 and shape.top + shape.height <= height + 10000 for slide in presentation.slides for shape in slide.shapes)
    elif item["kind"] == "image":
        from PIL import Image, ImageChops
        with Image.open(path) as picture:
            result["dimensions"], result["format"] = picture.size, picture.format
            rgb = picture.convert("RGB")
            red, green, blue = rgb.split()
            first = ImageChops.difference(red, green).getextrema()
            second = ImageChops.difference(red, blue).getextrema()
            assert isinstance(first[1], (int, float)) and isinstance(second[1], (int, float))
            result["grayscale_channel_delta"] = max(first[1], second[1])
            result["grayscale"] = result["grayscale_channel_delta"] <= 1
    elif path.suffix == ".docx":
        from docx import Document
        document = Document(str(path))
        result["paragraphs"] = [p.text for p in document.paragraphs if p.text]
        result["tables"] = len(document.tables)
    elif path.suffix == ".pdf":
        from pypdf import PdfReader
        pdf = PdfReader(path)
        result["pages"] = len(pdf.pages)
        result["text"] = "\n".join(page.extract_text() or "" for page in pdf.pages)
        import pypdfium2 as pdfium
        rendered = pdfium.PdfDocument(path)
        try:
            for number in range(min(2, len(rendered))):
                page = rendered[number]
                bitmap = cast(Any, page).render(scale=1.3)
                bitmap.to_pil().save(root / f"logs/quality/native/document-page-{number + 1}.png")
                bitmap.close()
                page.close()
        finally:
            rendered.close()
    elif item["kind"] == "audio":
        with wave.open(str(path)) as voice:
            result["duration"] = voice.getnframes() / voice.getframerate()
    results.append(result)
db.close()
target = root / "logs/quality/native/designed-artifact-content.json"
target.write_text(json.dumps(results, indent=2), encoding="utf-8")
print(json.dumps({"artifacts": len(results), "report": str(target)}))
